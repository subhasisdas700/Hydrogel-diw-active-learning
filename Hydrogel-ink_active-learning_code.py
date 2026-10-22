from __future__ import annotations
import argparse
from contextlib import contextmanager
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import unicodedata
import warnings

import joblib
import numpy as np
import pandas as pd
import scipy
from scipy.stats import norm
import sklearn
from sklearn.ensemble import RandomForestRegressor
from sklearn.gaussian_process import GaussianProcessRegressor
from sklearn.gaussian_process.kernels import ConstantKernel, RBF
from sklearn.metrics import mean_squared_error, r2_score
from sklearn.preprocessing import StandardScaler
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = "1.1.0-reconstructed"
INPUT_DATA_BASENAME = "Input data file-1.1"
DEFAULT_RUN_DIRECTORY = "al_run"
ID = "Sample-ID"
SCORE = "Enhanced_Score"
PRESSURE = "ΔP_used (MPa)"
COMPOSITION_COLUMNS = ["Alginate_wt%", "PVA_wt%", "Gelatin_wt%", "Glycerol_wt%"]
DERIVED_COLUMNS = ["Total_polymer_wt%", "Polymer_over_Glycerol"]
RHEOLOGY_COLUMNS = ["τy (Pa)", "τyd (Pa)", "n", "K", "G' at LVR", "G'' at LVR", "Loss Factor"]
COMPOSITION_FEATURES = COMPOSITION_COLUMNS + DERIVED_COLUMNS
FEATURE_COLUMNS = COMPOSITION_FEATURES + RHEOLOGY_COLUMNS
MEASUREMENT_COLUMNS = [ID] + COMPOSITION_COLUMNS + RHEOLOGY_COLUMNS + [SCORE, PRESSURE]
PROPOSAL_COLUMNS = COMPOSITION_COLUMNS + RHEOLOGY_COLUMNS + DERIVED_COLUMNS + [
    "Estimated_ΔP_HB (MPa)", "EI_Score", "Predicted_Score", "Predicted_Std",
    "Predicted_ΔP_used (MPa)", "P_used_over_HB_ratio", ID, "Water_wt%",
]
REPORT_COLUMNS = [ID, "Predicted_Enhanced_Score", "Predicted_Std", "Observed_Enhanced_Score"]


def default_config() -> dict:
    """Return independent, JSON-serializable settings; dimensions are explicit."""
    return {
        "initial_count": 25,
        "batch_size": 2,
        "pool_size": 2000,
        "xi": 0.01,
        "alpha": 0.01,
        "optimizer_restarts": 6,
        "rf_trees": 200,
        "random_seed": 42,
        "max_iterations": 10,
        "composition_match_atol": 1e-6,
        "geometry": {
            "nozzle_radius_mm": 0.30,
            "nozzle_length_mm": 19.7,
            "line_width_mm": 0.60,
            "line_height_mm": None,
            "print_speed_mm_s": 20.0,
        },
    }


def load_config(path: str | Path | None = None, batch_size: int | None = None) -> dict:
    config = default_config()
    if path:
        supplied = json.loads(Path(path).read_text(encoding="utf-8"))
        if not isinstance(supplied, dict):
            raise ValueError("The configuration must be a JSON object.")
        unknown = set(supplied) - set(config)
        if unknown:
            raise ValueError(f"Unknown configuration keys: {sorted(unknown)}")
        geometry = supplied.pop("geometry", {})
        if not isinstance(geometry, dict) or set(geometry) - set(config["geometry"]):
            raise ValueError("Unknown geometry settings; use the documented millimetre-based keys.")
        config.update(supplied)
        config["geometry"].update(geometry)
    if batch_size is not None:
        config["batch_size"] = batch_size
    validate_config(config)
    return config


def validate_config(config: dict) -> None:
    for key in ["initial_count", "batch_size", "pool_size", "rf_trees", "max_iterations"]:
        if isinstance(config[key], bool) or not isinstance(config[key], int) or config[key] < 1:
            raise ValueError(f"{key} must be a positive integer.")
    if config["initial_count"] < 8:
        raise ValueError("At least eight initial measured samples are required for the rheology regressors.")
    if config["batch_size"] > config["pool_size"]:
        raise ValueError("batch_size cannot exceed pool_size.")
    for key in ["optimizer_restarts", "random_seed"]:
        if isinstance(config[key], bool) or not isinstance(config[key], int) or config[key] < 0:
            raise ValueError(f"{key} must be a nonnegative integer.")
    for key in ["xi", "alpha", "composition_match_atol"]:
        if isinstance(config[key], bool) or not isinstance(config[key], (float, int)) or not np.isfinite(config[key]) or config[key] < 0:
            raise ValueError(f"{key} must be a finite nonnegative number.")
    if config["alpha"] == 0:
        raise ValueError("alpha must be positive for this noisy-measurement workflow.")
    for key, value in config["geometry"].items():
        if key == "line_height_mm" and value is None:
            continue
        if not isinstance(value, (float, int)) or not np.isfinite(value) or value <= 0:
            raise ValueError(f"geometry.{key} must be a positive number, or line_height_mm may be null.")


def normalized_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    text = text.replace("τ", "tau").replace("δ", "delta")
    text = text.replace("″", "''").replace("′", "'").replace("’", "'").replace("‘", "'")
    text = text.replace('g"', "gdoubleprime").replace("g''", "gdoubleprime").replace("g'", "gprime")
    return re.sub(r"[^a-z0-9]", "", text)


_ALIASES = {
    ID: ["Sample ID", "Sample_ID", "sample", "ID"],
    "Alginate_wt%": ["Alginate", "Sodium alginate wt%", "Sodium alginate"],
    "PVA_wt%": ["PVA", "PVA wt%"],
    "Gelatin_wt%": ["Gelatin", "Gelatin wt%"],
    "Glycerol_wt%": ["Glycerol", "Glycerol wt%"],
    "τy (Pa)": ["?y (Pa)", "tau_y (Pa)", "Yield stress (Pa)", "Yield stress", "tauy"],
    "τyd (Pa)": ["?yd (Pa)", "tau_yd (Pa)", "Dynamic yield stress (Pa)", "tauyd"],
    "n": ["Flow behaviour index", "Flow behavior index"],
    "K": ["Consistency index", "Consistency index K", "K (Pa.s^n)", "K (Pa·sⁿ)"],
    "G' at LVR": ["G′ at LVR", "G' at LVR (Pa)", "Storage modulus (Pa)", "G'", "Gprime"],
    "G'' at LVR": ["G″ at LVR", 'G" at LVR', "G'' at LVR (Pa)", "Loss modulus (Pa)", "G''", "Gdoubleprime"],
    "Loss Factor": ["Loss_factor", "tan delta", "tan δ"],
    SCORE: ["Enhanced Score", "Total printability Score", "Printability score", "Observed_Enhanced_Score"],
    PRESSURE: ["?P_used (MPa)", "P_used (MPa)", "Applied pressure (MPa)", "Extrusion pressure (MPa)", "DeltaP_used (MPa)"],
}
_HEADER_LOOKUP = {normalized_header(alias): key for key, names in _ALIASES.items() for alias in [key] + names}


def resolve_input_file(path: str | Path | None = None) -> Path:
    """Resolve the complete original basename, preserving the '.1' in its name."""
    base = Path(INPUT_DATA_BASENAME if path is None else path)
    if base.is_file():
        if base.suffix.casefold() not in {".xlsx", ".csv"}:
            raise ValueError("The input file must have a .xlsx or .csv extension.")
        return base.resolve()
    if base.suffix.casefold() in {".xlsx", ".csv"}:
        raise FileNotFoundError(f"Input file not found: {base}")
    matches = [Path(str(base) + extension) for extension in [".xlsx", ".csv"]
               if Path(str(base) + extension).is_file()]
    if len(matches) > 1:
        raise ValueError(f"Both Excel and CSV versions of {base.name!r} exist. Select one explicitly with --data and its full filename.")
    if not matches:
        raise FileNotFoundError(f"Cannot find {base}.xlsx or {base}.csv. Run from the input file's folder or specify --data with its full path.")
    return matches[0].resolve()


def read_measurements(path: str | Path, sheet: str | int = 0) -> pd.DataFrame:
    path = Path(path)
    if path.suffix.casefold() == ".csv":
        try:
            frame = pd.read_csv(path, encoding="utf-8-sig")
        except UnicodeDecodeError:
            frame = pd.read_csv(path, encoding="cp1252")
    elif path.suffix.casefold() == ".xlsx":
        frame = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
    else:
        raise ValueError("Input must be .xlsx or .csv; save older .xls workbooks as .xlsx first.")
    return frame.dropna(how="all").reset_index(drop=True)


def add_derived(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    result["Total_polymer_wt%"] = result[COMPOSITION_COLUMNS[:3]].sum(axis=1)
    result["Polymer_over_Glycerol"] = result["Total_polymer_wt%"] / result["Glycerol_wt%"]
    result["Water_wt%"] = 100.0 - result[COMPOSITION_COLUMNS].sum(axis=1)
    return result


def prepare_measured_data(frame: pd.DataFrame, generate_ids: bool = False) -> pd.DataFrame:
    """Validate real observations; never fill unknown measurements with predictions."""
    frame = frame.dropna(how="all").reset_index(drop=True)
    result = pd.DataFrame(index=frame.index)
    for column in frame.columns:
        canonical = _HEADER_LOOKUP.get(normalized_header(column))
        if canonical is None:
            continue
        values = frame[column]
        if canonical in result:
            a, b = result[canonical], values
            if canonical == ID:
                agree = a.fillna("").astype(str).str.strip().eq(b.fillna("").astype(str).str.strip()).all()
            else:
                av, bv = pd.to_numeric(a, errors="coerce"), pd.to_numeric(b, errors="coerce")
                agree = np.allclose(av, bv, equal_nan=True, rtol=1e-10, atol=1e-12)
            if not agree:
                raise ValueError(f"Conflicting columns map to {canonical!r}; keep one authoritative column.")
        else:
            result[canonical] = values
    if ID not in result:
        if not generate_ids:
            raise ValueError("Sample-ID is required for batch measurements; retain the IDs in the generated template.")
        result[ID] = [f"S{i + 1}" for i in range(len(result))]
    if result[ID].isna().any():
        raise ValueError("Sample-ID contains a blank cell.")
    result[ID] = result[ID].astype(str).str.strip()
    if result[ID].eq("").any() or result[ID].str.casefold().isin(["nan", "none"]).any():
        raise ValueError("Sample-ID contains an empty identifier.")
    if result[ID].duplicated().any():
        raise ValueError(f"Duplicate Sample-ID values: {result.loc[result[ID].duplicated(), ID].tolist()}")
    required = COMPOSITION_COLUMNS + RHEOLOGY_COLUMNS[:-1] + [SCORE]
    missing = [column for column in required if column not in result]
    if missing:
        raise ValueError(f"Missing measured columns: {missing}. Use the supplied input template.")
    for column in required + ["Loss Factor", PRESSURE]:
        if column not in result:
            result[column] = np.nan
        original = result[column]
        result[column] = pd.to_numeric(original, errors="coerce")
        invalid_text = original.notna() & result[column].isna() & original.astype(str).str.strip().ne("")
        if invalid_text.any():
            raise ValueError(f"Nonnumeric data in {column!r}, sample(s): {result.loc[invalid_text, ID].tolist()}")
        if np.isinf(result[column].to_numpy(dtype=float)).any():
            raise ValueError(f"Infinite value in {column!r}.")
    for column in required:
        if result[column].isna().any():
            raise ValueError(f"Missing real measurement in {column!r}, sample(s): {result.loc[result[column].isna(), ID].tolist()}")
    result["Loss Factor"] = result["Loss Factor"].fillna(result["G'' at LVR"] / result["G' at LVR"])
    if not np.isfinite(result[COMPOSITION_COLUMNS + RHEOLOGY_COLUMNS + [SCORE]].to_numpy(dtype=float)).all():
        raise ValueError("Required observations must all be finite, including the calculated loss factor.")
    if (result[COMPOSITION_COLUMNS] < 0).any().any() or (result["Glycerol_wt%"] <= 0).any():
        raise ValueError("Compositions must be nonnegative and glycerol must be positive for the polymer:glycerol descriptor.")
    if (result[COMPOSITION_COLUMNS].sum(axis=1) > 100.0 + 1e-9).any():
        raise ValueError("The four component weight percentages exceed 100% in at least one formulation.")
    for column in ["τy (Pa)", "τyd (Pa)", "G'' at LVR", "Loss Factor"]:
        if (result[column] < 0).any():
            raise ValueError(f"{column} must be nonnegative.")
    for column in ["n", "K", "G' at LVR"]:
        if (result[column] <= 0).any():
            raise ValueError(f"{column} must be positive.")
    if ((result[SCORE] < 0) | (result[SCORE] > 1)).any():
        raise ValueError("Measured Enhanced_Score must already be on [0, 1]; scores are never renormalized between iterations.")
    if (result[PRESSURE].dropna() < 0).any():
        raise ValueError("Measured applied pressures must be nonnegative MPa values.")
    return add_derived(result[MEASUREMENT_COLUMNS])


def make_gpr(config: dict, alpha: float | None = None) -> GaussianProcessRegressor:
    kernel = ConstantKernel(1.0, (1e-3, 1e3)) * RBF(1.0, (1e-2, 1e2))
    return GaussianProcessRegressor(kernel=kernel, alpha=config["alpha"] if alpha is None else alpha,
                                    n_restarts_optimizer=config["optimizer_restarts"], normalize_y=False,
                                    random_state=config["random_seed"])


def fit_models(measured: pd.DataFrame, config: dict, scaler: StandardScaler | None = None) -> dict:
    """Fit only observed rows; passing a scaler preserves the initial coordinate system."""
    features = measured[FEATURE_COLUMNS].astype(float)
    if scaler is None:
        if len(measured) != config["initial_count"]:
            raise ValueError("A fresh scaler must be fitted on exactly initial_count observed rows.")
        scaler = StandardScaler().fit(features)
    scaled = scaler.transform(features)
    gpr = make_gpr(config)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        gpr.fit(scaled, measured[SCORE].to_numpy(dtype=float))
        rheology = {}
        for column in RHEOLOGY_COLUMNS:
            rf = RandomForestRegressor(n_estimators=config["rf_trees"], random_state=config["random_seed"], n_jobs=-1)
            rf.fit(measured[COMPOSITION_FEATURES], measured[column])
            rheology[column] = rf
        pressure = None
        pressure_rows = measured[PRESSURE].notna().to_numpy()
        if int(pressure_rows.sum()) >= 8:
            pressure = make_gpr(config, alpha=1e-3)
            pressure.fit(scaled[pressure_rows], measured.loc[pressure_rows, PRESSURE].to_numpy(dtype=float))
    mu, sigma = gpr.predict(scaled, return_std=True)
    return {"gpr": gpr, "scaler": scaler, "rheology_models": rheology, "pressure_model": pressure,
            "measured": measured.copy(), "config": deepcopy(config),
            "training_metrics": prediction_metrics(measured[SCORE].to_numpy(), mu, sigma),
            "fit_warnings": list(dict.fromkeys(str(w.message) for w in caught))}


def expected_improvement(mu, sigma, best_observed: float, xi: float = 0.01) -> np.ndarray:
    mu, sigma = np.broadcast_arrays(np.asarray(mu, dtype=float), np.asarray(sigma, dtype=float))
    if not np.isfinite(mu).all() or not np.isfinite(sigma).all() or (sigma < 0).any():
        raise ValueError("EI requires finite predictive means and nonnegative finite standard deviations.")
    improvement = mu - float(best_observed) - float(xi)
    result = np.array(np.maximum(improvement, 0.0), copy=True)
    positive = sigma > 0
    z = improvement[positive] / sigma[positive]
    result[positive] = improvement[positive] * norm.cdf(z) + sigma[positive] * norm.pdf(z)
    return np.maximum(result, 0.0)


def hb_pressure(frame: pd.DataFrame, geometry: dict) -> np.ndarray:
    """Approximate HB nozzle pressure in MPa; blank when filament height is unknown."""
    if geometry["line_height_mm"] is None:
        return np.full(len(frame), np.nan)
    radius = geometry["nozzle_radius_mm"] * 1e-3
    length = geometry["nozzle_length_mm"] * 1e-3
    area = geometry["line_width_mm"] * geometry["line_height_mm"] * 1e-6
    flow = geometry["print_speed_mm_s"] * 1e-3 * area
    gamma_app = 4.0 * flow / (np.pi * radius ** 3)
    n = frame["n"].to_numpy(dtype=float)
    gamma_wall = (3.0 * n + 1.0) / (4.0 * n) * gamma_app
    tau_wall = frame["τy (Pa)"].to_numpy() + frame["K"].to_numpy() * gamma_wall ** n
    return 2.0 * length / radius * tau_wall / 1e6


def _next_ids(existing: list[str], count: int) -> list[str]:
    occupied = set(existing)
    numeric = [int(match.group(1)) for value in existing if (match := re.fullmatch(r"S(\d+)", value))]
    number = max(numeric, default=len(existing)) + 1
    result = []
    while len(result) < count:
        candidate = f"S{number}"
        if candidate not in occupied:
            result.append(candidate)
        number += 1
    return result


def propose_candidates(models: dict, iteration: int, batch_size: int | None = None) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    config, measured = models["config"], models["measured"]
    batch_size = config["batch_size"] if batch_size is None else batch_size
    if not isinstance(batch_size, int) or batch_size < 1 or batch_size > config["pool_size"]:
        raise ValueError("batch_size must be between 1 and pool_size.")
    lower = measured[COMPOSITION_COLUMNS].min().to_numpy(dtype=float)
    upper = measured[COMPOSITION_COLUMNS].max().to_numpy(dtype=float)
    if np.array_equal(lower, upper):
        raise ValueError("All observed compositions are identical; observed min–max bounds contain no new candidates.")
    seed = int(config["random_seed"] + iteration - 1)
    rng = np.random.default_rng(seed)
    occupied = {tuple(np.round(row, 10)) for row in measured[COMPOSITION_COLUMNS].to_numpy(dtype=float)}
    accepted = []
    for _ in range(100):
        draw = rng.uniform(lower, upper, size=(config["pool_size"] - len(accepted), 4))
        for row in draw:
            key = tuple(np.round(row, 10))
            if row.sum() <= 100.0 and row[3] > 0 and key not in occupied:
                accepted.append(row)
                occupied.add(key)
        if len(accepted) == config["pool_size"]:
            break
    if len(accepted) < config["pool_size"]:
        raise ValueError("Cannot draw the requested number of distinct valid formulations from the observed bounds.")
    pool = add_derived(pd.DataFrame(accepted, columns=COMPOSITION_COLUMNS))
    for column, model in models["rheology_models"].items():
        pool[column] = model.predict(pool[COMPOSITION_FEATURES])
    scaled = models["scaler"].transform(pool[FEATURE_COLUMNS])
    means, deviations = [], []
    for start in range(0, len(pool), 2000):
        mean, deviation = models["gpr"].predict(scaled[start:start + 2000], return_std=True)
        means.append(mean)
        deviations.append(deviation)
    pool["Predicted_Score"] = np.concatenate(means)
    pool["Predicted_Std"] = np.concatenate(deviations)
    pool["EI_Score"] = expected_improvement(pool["Predicted_Score"], pool["Predicted_Std"], measured[SCORE].max(), config["xi"])
    pool["Estimated_ΔP_HB (MPa)"] = hb_pressure(pool, config["geometry"])
    pressure_model = models["pressure_model"]
    pool["Predicted_ΔP_used (MPa)"] = pressure_model.predict(scaled) if pressure_model is not None else np.nan
    pool["P_used_over_HB_ratio"] = pool["Predicted_ΔP_used (MPa)"] / pool["Estimated_ΔP_HB (MPa)"].replace(0, np.nan)
    pool[ID] = ""
    selected_indices = pool["EI_Score"].sort_values(ascending=False, kind="stable").index[:batch_size]
    selected = pool.loc[selected_indices].copy().reset_index(drop=True)
    selected[ID] = _next_ids(measured[ID].tolist(), batch_size)
    pool.loc[selected_indices, ID] = selected[ID].to_numpy()
    info = {"iteration": int(iteration), "candidate_seed": seed, "pool_size": len(pool), "batch_size": len(selected),
            "best_observed_score": float(measured[SCORE].max()), "xi": float(config["xi"]),
            "composition_bounds": {name: [float(low), float(high)] for name, low, high in zip(COMPOSITION_COLUMNS, lower, upper)},
            "selection": "Top-k raw expected improvement; no OOD, pressure, or adjusted-score filtering.",
            "pressure_is_informational_only": True,
            "negative_applied_pressure_predictions": int((pool["Predicted_ΔP_used (MPa)"] < 0).sum()),
            "pressure_notes": [note for condition, note in [
                (config["geometry"]["line_height_mm"] is None,
                 "Estimated HB pressure and pressure ratio are blank: filament height is unspecified."),
                (pressure_model is None,
                 "Predicted applied pressure is blank: fewer than eight measured applied-pressure values."),
                (pressure_model is not None,
                 "Applied-pressure GP predictions are raw regression estimates and may be negative; they are not pressure settings or feasibility approval."),
                (bool((pool["Predicted_ΔP_used (MPa)"] < 0).any()),
                 "This pool contains negative applied-pressure predictions; these are recorded without clipping and do not affect EI ranking."),
            ] if condition]}
    return selected[PROPOSAL_COLUMNS], pool[PROPOSAL_COLUMNS], info


def prediction_metrics(observed, predicted, sigma) -> dict:
    observed, predicted, sigma = map(lambda x: np.asarray(x, dtype=float), (observed, predicted, sigma))
    n = len(observed)
    return {"n": int(n), "rmse": float(np.sqrt(mean_squared_error(observed, predicted))),
            "mae": float(np.mean(np.abs(observed - predicted))),
            "r2": float(r2_score(observed, predicted)) if n >= 3 and np.ptp(observed) > 0 else None,
            "mean_gp_std": float(sigma.mean())}


def report_frame(selected: pd.DataFrame, observed: pd.DataFrame | None = None) -> pd.DataFrame:
    result = pd.DataFrame({ID: selected[ID].to_numpy(), "Predicted_Enhanced_Score": selected["Predicted_Score"].to_numpy(),
                           "Predicted_Std": selected["Predicted_Std"].to_numpy(), "Observed_Enhanced_Score": np.nan})
    if observed is not None:
        result["Observed_Enhanced_Score"] = observed.set_index(ID).loc[result[ID], SCORE].to_numpy()
    return result[REPORT_COLUMNS]


def utc_timestamp() -> str:
    return datetime.now(timezone.utc).isoformat()


def file_hash(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def versions() -> dict:
    import openpyxl
    return {"python": sys.version.split()[0], "numpy": np.__version__, "pandas": pd.__version__,
            "scipy": scipy.__version__, "scikit-learn": sklearn.__version__,
            "openpyxl": openpyxl.__version__, "joblib": joblib.__version__}


def write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    """Write simple, readable Python-generated Excel outputs without hidden formulas."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.book[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 60
            for cell in sheet[1]:
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="163D50")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for index, column in enumerate(frame.columns, 1):
                width = 16 if column != ID else 15
                if name == "Notes":
                    width = 30 if index == 1 else 100
                elif len(str(column)) > 24:
                    width = 23
                sheet.column_dimensions[get_column_letter(index)].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Calibri", size=11)
                    cell.alignment = Alignment(vertical="center", wrap_text=name == "Notes")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.000000"
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="EEF4F7")
                sheet.row_dimensions[row[0].row].height = 44 if name == "Notes" else 22


def notes_frame(config: dict, extra: list[tuple[str, str]] | None = None) -> pd.DataFrame:
    notes = [
        ("Implementation", "Reconstructed from SI S.3/S.4 and examples; not the recovered original code."),
        ("Proposed candidate properties", "The rheological properties in Proposed candidates are RF estimates, not experimental measurements."),
        ("Prediction source", "The main prediction report preserves the GP predictions made at proposal time, using RF-estimated rheology."),
        ("Observed score", "Blank until real experiments are supplied. New observed scores are never used to generate their own prior predictions."),
        ("GP standard deviation", "Conditional GP uncertainty; does not propagate uncertainty in RF-predicted rheology or include all experimental uncertainty."),
        ("Bounds and ranking", "2,000 candidates by default, uniformly sampled within the current measured min–max composition bounds; rank by EI."),
        ("Pressure", "Pressure columns are informational and do not filter active-learning candidates. Missing pressure predictions remain blank."),
        ("HB geometry", "HB pressure is blank until line_height_mm is explicitly supplied in the configuration." if config["geometry"]["line_height_mm"] is None else json.dumps(config["geometry"])),
        ("GPR scores", "Raw GP means are exported and used in EI; predictions can lie outside [0,1]. Measured input scores must lie in [0,1]."),
    ]
    return pd.DataFrame(notes + (extra or []), columns=["Item", "Meaning"])


def save_model_snapshot(run: Path, models: dict, iteration: int) -> str:
    directory = run / "models"
    directory.mkdir(exist_ok=True)
    name = f"models/model_after_iteration_{iteration:02d}.joblib"
    joblib.dump(models, run / name, compress=3)
    measured_file = run / f"measured_after_iteration_{iteration:02d}.csv"
    models["measured"].to_csv(measured_file, index=False, encoding="utf-8-sig")
    summary = {"completed_iteration": iteration, "training_rows": len(models["measured"]),
               "initial_scaler_rows": int(models["scaler"].n_samples_seen_),
               "fitted_gp_kernel": str(models["gpr"].kernel_),
               "training_metrics_only": models["training_metrics"], "fit_warnings": models["fit_warnings"],
               "measured_snapshot_sha256": file_hash(measured_file), "versions": versions()}
    write_json(run / f"model_summary_iteration_{iteration:02d}.json", summary)
    return name


def save_proposal(run: Path, models: dict, iteration: int, batch_size: int) -> None:
    directory = run / f"iteration_{iteration:02d}"
    if directory.exists():
        raise FileExistsError(f"Iteration output already exists: {directory}")
    directory.mkdir()
    selected, pool, info = propose_candidates(models, iteration, batch_size)
    joblib.dump({"selected": selected, "info": info}, directory / "proposal_snapshot.joblib", compress=3)
    pool.to_csv(directory / f"candidate_pool_iteration_{iteration:02d}.csv", index=False, encoding="utf-8-sig")
    notes = notes_frame(models["config"])
    write_workbook(directory / f"proposed_candidates_iteration_{iteration:02d}.xlsx", {"Proposed candidates": selected, "Notes": notes})
    write_workbook(directory / f"prediction_on_new_iteration_{iteration:02d}.xlsx", {"Predictions": report_frame(selected), "Notes": notes})
    measurements = pd.DataFrame(index=selected.index, columns=MEASUREMENT_COLUMNS)
    measurements[[ID] + COMPOSITION_COLUMNS] = selected[[ID] + COMPOSITION_COLUMNS]
    write_workbook(directory / f"measurements_iteration_{iteration:02d}.xlsx", {
        "Measurements": measurements,
        "Notes": notes_frame(models["config"], [("Complete this sheet", "Enter actual measured rheology and Enhanced_Score. Keep Sample-ID and compositions. Applied pressure is optional; Loss Factor can be left blank to calculate G''/G'.")]),
    })
    write_json(directory / "proposal_metadata.json", {**info, "created_utc": utc_timestamp(),
               "training_rows_at_proposal": len(models["measured"]), "fitted_gp_kernel": str(models["gpr"].kernel_),
               "config": models["config"], "versions": versions()})


def _run_state(config: dict, model_path: str, completed: int, pending: int | None, previous: dict | None = None) -> dict:
    return {**(previous or {}), "schema_version": 1, "implementation_version": VERSION,
            "config": deepcopy(config), "current_model": model_path, "completed_iteration": completed,
            "pending_iteration": pending, "last_updated_utc": utc_timestamp()}


def start_run(data_path: str | Path, output: str | Path, config: dict, sheet: str | int = 0) -> Path:
    output = Path(output).resolve()
    if output.exists():
        raise FileExistsError(f"Output path already exists: {output}. Choose a new run directory to preserve existing results.")
    validate_config(config)
    measured = prepare_measured_data(read_measurements(data_path, sheet), generate_ids=True)
    if len(measured) != config["initial_count"]:
        raise ValueError(f"Expected exactly {config['initial_count']} initial measured rows, found {len(measured)}. Start with the initial dataset only.")
    output.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=f".{output.name}.staging_", dir=output.parent))
    try:
        models = fit_models(measured, config)
        model_path = save_model_snapshot(stage, models, 0)
        save_proposal(stage, models, 1, config["batch_size"])
        state = _run_state(config, model_path, 0, 1)
        state.update({"created_utc": utc_timestamp(), "initial_sample_ids": measured[ID].tolist(),
                      "initial_source_name": Path(data_path).name, "initial_source_sha256": file_hash(data_path),
                      "script_sha256": file_hash(__file__), "versions": versions()})
        write_json(stage / "run_state.json", state)
        os.replace(stage, output)
    finally:
        if stage.exists():
            shutil.rmtree(stage)
    return output


@contextmanager
def run_lock(run: Path):
    # A sibling lock survives the transactional replacement of the run directory.
    lock = run.parent / f".{run.name}.update.lock"
    try:
        descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as error:
        raise RuntimeError(f"This run is already being updated. If an earlier process crashed, confirm it has stopped before removing {lock}.") from error
    os.close(descriptor)
    try:
        yield
    finally:
        lock.unlink(missing_ok=True)


def extract_new_master_rows(master: pd.DataFrame, previous: pd.DataFrame) -> pd.DataFrame:
    """Keep earlier measured rows immutable; select appended samples by stable ID."""
    old_ids = set(previous[ID])
    missing = sorted(old_ids - set(master[ID]))
    if missing:
        raise ValueError(f"The cumulative input is missing previously measured Sample-IDs: {missing}. Keep all earlier rows in the master file.")
    numeric = [column for column in MEASUREMENT_COLUMNS if column != ID]
    prior = previous.set_index(ID)
    existing = master.set_index(ID).loc[prior.index]
    equal = np.isclose(existing[numeric].to_numpy(dtype=float), prior[numeric].to_numpy(dtype=float),
                       rtol=1e-9, atol=1e-8, equal_nan=True)
    if not equal.all():
        changed = prior.index[~equal.all(axis=1)].tolist()
        raise ValueError(f"Previously recorded measurements changed for Sample-IDs: {changed}. This update accepts appended experiments; review corrections separately instead of silently changing the saved history.")
    return master.loc[~master[ID].isin(old_ids)].copy().reset_index(drop=True)


def update_run(run: str | Path, measurements_path: str | Path, batch_size: int | None = None,
               no_propose: bool = False, sheet: str | int = 0, cumulative: bool = False) -> Path:
    run = Path(run).resolve()
    if not (run / "run_state.json").is_file():
        raise ValueError("The run directory is missing run_state.json; start a run first.")
    with run_lock(run):
        state = json.loads((run / "run_state.json").read_text(encoding="utf-8"))
        iteration = state["pending_iteration"]
        if iteration is None:
            raise ValueError("This run has no pending batch. It was completed or stopped with --no-propose.")
        config = deepcopy(state["config"])
        if batch_size is not None:
            config["batch_size"] = batch_size
        validate_config(config)
        old_models = joblib.load(run / state["current_model"])
        proposal = joblib.load(run / f"iteration_{iteration:02d}" / "proposal_snapshot.joblib")
        selected = proposal["selected"]
        incoming = prepare_measured_data(read_measurements(measurements_path, sheet), generate_ids=cumulative)
        new = extract_new_master_rows(incoming, old_models["measured"]) if cumulative else incoming
        if cumulative and new.empty:
            print(f"No new measured samples in {Path(measurements_path).name}; iteration {iteration} is still awaiting its experimental batch. Saved results are unchanged.")
            return run
        if set(new[ID]) != set(selected[ID]):
            missing = sorted(set(selected[ID]) - set(new[ID]))
            extra = sorted(set(new[ID]) - set(selected[ID]))
            raise ValueError(f"Measurements must contain exactly the pending IDs. Missing: {missing}; extra: {extra}")
        new = new.set_index(ID).loc[selected[ID]].reset_index()
        close = np.isclose(new[COMPOSITION_COLUMNS].to_numpy(dtype=float), selected[COMPOSITION_COLUMNS].to_numpy(dtype=float),
                           atol=config["composition_match_atol"], rtol=0)
        if not close.all():
            bad = new.loc[~close.all(axis=1), ID].tolist()
            raise ValueError(f"Measured compositions do not match their proposed Sample-IDs: {bad}. Retain exact compositions or explicitly set an appropriate composition_match_atol before starting.")
        if set(new[ID]) & set(old_models["measured"][ID]):
            raise ValueError("A pending Sample-ID is already in the measured training data.")
        # These comparisons are deliberately constructed BEFORE fitting on new scores.
        main_report = report_frame(selected, new)
        conditional_mu, conditional_std = old_models["gpr"].predict(old_models["scaler"].transform(new[FEATURE_COLUMNS]), return_std=True)
        conditional_report = pd.DataFrame({ID: new[ID], "Predicted_Enhanced_Score": conditional_mu,
                                           "Predicted_Std": conditional_std, "Observed_Enhanced_Score": new[SCORE]})
        new_metrics = {"iteration": iteration, "new_batch_size": len(new),
                       "proposal_time_predictions": prediction_metrics(new[SCORE], selected["Predicted_Score"], selected["Predicted_Std"]),
                       "prior_gp_with_new_measured_rheology": prediction_metrics(new[SCORE], conditional_mu, conditional_std),
                       "interpretation": "Both comparisons precede refitting. The second conditions on newly measured rheology and is not the original composition-only forecast. Small-batch metrics are descriptive, not independent validation of the complete study.",
                       "measurement_source_name": Path(measurements_path).name,
                       "measurement_source_sha256": file_hash(measurements_path),
                       "input_mode": "cumulative_master" if cumulative else "separate_batch"}
        augmented = pd.concat([old_models["measured"], new], ignore_index=True)
        models = fit_models(augmented, config, scaler=old_models["scaler"])
        if int(models["scaler"].n_samples_seen_) != config["initial_count"]:
            raise RuntimeError("The initial scaler was unexpectedly altered; update cancelled.")
        stage = Path(tempfile.mkdtemp(prefix=f".{run.name}.staging_", dir=run.parent))
        backup = None
        try:
            shutil.copytree(run, stage, dirs_exist_ok=True, ignore=shutil.ignore_patterns(".update.lock"))
            directory = stage / f"iteration_{iteration:02d}"
            extra = [("Measured rheology prediction sheet", "Uses newly measured rheology with the prior GP. The new scores were excluded from that GP; these conditional predictions are different from the saved proposal-time predictions."),
                     ("Batch metrics", "Descriptive diagnostics only. R² is omitted for batches smaller than three or constant observed scores.")]
            write_workbook(directory / f"prediction_on_new_iteration_{iteration:02d}.xlsx", {
                "Predictions": main_report, "Measured rheology prediction": conditional_report,
                "Notes": notes_frame(config, extra),
            })
            new.to_csv(directory / "observed_measurements.csv", index=False, encoding="utf-8-sig")
            write_json(directory / "new_batch_metrics.json", new_metrics)
            model_path = save_model_snapshot(stage, models, iteration)
            pending = None
            if not no_propose and iteration < config["max_iterations"]:
                pending = iteration + 1
                save_proposal(stage, models, pending, config["batch_size"])
            updated = _run_state(config, model_path, iteration, pending, state)
            updated["last_update_script_sha256"] = file_hash(__file__)
            write_json(stage / "run_state.json", updated)
            # Build everything before replacing the run, keeping a recoverable backup.
            backup = Path(tempfile.mkdtemp(prefix=f".{run.name}.backup_", dir=run.parent))
            backup.rmdir()
            os.replace(run, backup)
            try:
                os.replace(stage, run)
            except BaseException:
                os.replace(backup, run)
                backup = None
                raise
            shutil.rmtree(backup)
            backup = None
        finally:
            if stage.exists():
                shutil.rmtree(stage)
        return run


def run_master_file(data_path: str | Path | None = None, output: str | Path = DEFAULT_RUN_DIRECTORY,
                    config_path: str | Path | None = None, batch_size: int | None = None,
                    no_propose: bool = False, sheet: str | int = 0) -> Path:
    """One command for the user's continually updated master input file."""
    source = resolve_input_file(data_path)
    run = Path(output).resolve()
    print(f"Input data: {source}")
    if (run / "run_state.json").is_file():
        if config_path is not None:
            raise ValueError("This run already has saved settings. Omit --config when updating; --batch-size can change the next batch.")
        return update_run(run, source, batch_size, no_propose, sheet, cumulative=True)
    if no_propose:
        raise ValueError("--no-propose applies when accepting an existing pending batch, not when starting.")
    if config_path is None and Path("active_learning_config.json").is_file():
        config_path = Path("active_learning_config.json")
    return start_run(source, run, load_config(config_path, batch_size), sheet)


def print_run_summary(run: Path) -> None:
    state = json.loads((run / "run_state.json").read_text(encoding="utf-8"))
    print(f"Results: {run}")
    print(f"Completed experimental iterations: {state['completed_iteration']}")
    pending = state["pending_iteration"]
    if pending is not None:
        path = run / f"iteration_{pending:02d}"
        print(f"Proposed candidates: {path / f'proposed_candidates_iteration_{pending:02d}.xlsx'}")
        print("Append this batch's actual measurements to your master input file, retaining earlier rows and Sample-IDs.")
        print(f"Enter real measurements in a copy of: {path / f'measurements_iteration_{pending:02d}.xlsx'}")
        print(f"Prediction report: {path / f'prediction_on_new_iteration_{pending:02d}.xlsx'}")
    else:
        print("No further batch proposed. The measured dataset and final fitted model are saved.")
    if state["config"]["geometry"]["line_height_mm"] is None:
        print("HB pressure is blank: line_height_mm has not been specified in the configuration.")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    auto = commands.add_parser("run", help="Start or update using the same cumulative input file; also the no-argument default.")
    auto.add_argument("--data", help="Master .xlsx/.csv or basename; defaults to Input data file-1.1.")
    auto.add_argument("--output", default=DEFAULT_RUN_DIRECTORY, help="Persistent run directory; default al_run.")
    auto.add_argument("--config", help="JSON settings for the first run only; otherwise active_learning_config.json is used when present.")
    auto.add_argument("--batch-size", type=int, help="Number of candidates in the next proposal; default 2 on a new run.")
    auto.add_argument("--no-propose", action="store_true", help="Accept a measured batch and stop without another proposal.")
    auto.add_argument("--sheet", default=None, help="Worksheet name; default is the first sheet.")
    start = commands.add_parser("start", help="Train on the initial measured samples and propose the first batch.")
    start.add_argument("--data", help="Initial measured .xlsx or .csv; defaults to Input data file-1.1, exactly 25 rows.")
    start.add_argument("--output", default=DEFAULT_RUN_DIRECTORY, help="New run directory; an existing directory is never overwritten.")
    start.add_argument("--config", help="Optional JSON overrides; see example configuration.")
    start.add_argument("--batch-size", type=int, help="New formulations per batch; default 2, or use 4.")
    start.add_argument("--sheet", default=None, help="Input worksheet name; default is the first sheet.")
    update = commands.add_parser("update", help="Evaluate a measured batch, retrain, and propose the next batch.")
    update.add_argument("--run", default=DEFAULT_RUN_DIRECTORY, help="Run directory created by start; default al_run.")
    inputs = update.add_mutually_exclusive_group()
    inputs.add_argument("--data", help="Cumulative master file, including earlier and new measured rows. Default: Input data file-1.1.xlsx/.csv.")
    inputs.add_argument("--measurements", help="Alternative: separate file containing only the pending measured batch.")
    update.add_argument("--batch-size", type=int, help="Size of the NEXT proposed batch; pending batch size is unchanged.")
    update.add_argument("--no-propose", action="store_true", help="Evaluate and refit after this batch, then end this run without another proposal.")
    update.add_argument("--sheet", default=None, help="Input worksheet name; default is the first sheet.")
    argv = list(sys.argv[1:] if argv is None else argv)
    args = parser.parse_args(argv or ["run"])
    try:
        if args.command == "run":
            run = run_master_file(args.data, args.output, args.config, args.batch_size, args.no_propose, args.sheet or 0)
        elif args.command == "start":
            run = start_run(resolve_input_file(args.data), args.output, load_config(args.config, args.batch_size), args.sheet or 0)
        else:
            cumulative = args.measurements is None
            source = resolve_input_file(args.data) if cumulative else args.measurements
            run = update_run(args.run, source, args.batch_size, args.no_propose, args.sheet or 0, cumulative=cumulative)
        print_run_summary(run)
        return 0
    except (ValueError, OSError, RuntimeError, KeyError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
